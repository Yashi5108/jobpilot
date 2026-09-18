from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from backend.services.resume_parser import (
    ResumeParseError,
    normalize_resume_text,
    parse_resume,
)
from backend.services.resume_validation import (
    ResumeValidationError,
    validate_resume_upload,
)


def _escape_pdf_text(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _build_pdf_with_text(lines: list[str]) -> bytes:
    page_count = len(lines)
    objects: dict[int, bytes] = {}

    font_id = 3 + (page_count * 2)
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"

    kids_refs: list[str] = []
    for idx, text in enumerate(lines):
        page_obj_id = 3 + (idx * 2)
        content_obj_id = page_obj_id + 1
        kids_refs.append(f"{page_obj_id} 0 R")

        escaped = _escape_pdf_text(text)
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("utf-8")
        objects[content_obj_id] = (
            b"<< /Length "
            + str(len(stream)).encode("utf-8")
            + b" >>\n"
            + b"stream\n"
            + stream
            + b"\nendstream"
        )
        objects[page_obj_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            + f"/Contents {content_obj_id} 0 R ".encode("utf-8")
            + f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode("utf-8")
        )

    kids = " ".join(kids_refs)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode(
        "utf-8"
    )
    objects[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    max_id = max(objects)
    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}

    for obj_id in range(1, max_id + 1):
        offsets[obj_id] = len(output)
        output.extend(f"{obj_id} 0 obj\n".encode("utf-8"))
        output.extend(objects[obj_id])
        output.extend(b"\nendobj\n")

    xref_offset = len(output)
    output.extend(f"xref\n0 {max_id + 1}\n".encode("utf-8"))
    output.extend(b"0000000000 65535 f \n")
    for obj_id in range(1, max_id + 1):
        output.extend(f"{offsets[obj_id]:010d} 00000 n \n".encode("utf-8"))

    output.extend(
        (
            f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("utf-8")
    )
    return bytes(output)


def test_parse_pdf_extracts_text_and_pages(tmp_path: Path) -> None:
    pdf_path = tmp_path / "resume.pdf"
    pdf_path.write_bytes(_build_pdf_with_text(["Alice Resume", "Python Developer"]))

    parsed = parse_resume(pdf_path, "pdf", "resume.pdf")

    assert parsed.file_type == "pdf"
    assert parsed.page_count == 2
    assert "Alice Resume" in parsed.normalized_text
    assert "Python Developer" in parsed.normalized_text


def test_parse_pdf_empty_raises(tmp_path: Path) -> None:
    pdf_path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with pdf_path.open("wb") as stream:
        writer.write(stream)

    with pytest.raises(ResumeParseError, match="no extractable text"):
        parse_resume(pdf_path, "pdf", "empty.pdf")


def test_parse_pdf_corrupted_raises(tmp_path: Path) -> None:
    pdf_path = tmp_path / "broken.pdf"
    pdf_path.write_bytes(b"%PDF-not-valid")

    with pytest.raises(ResumeParseError, match="Invalid or corrupted PDF"):
        parse_resume(pdf_path, "pdf", "broken.pdf")


def test_parse_docx_extracts_paragraphs_and_tables(tmp_path: Path) -> None:
    docx_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Alex Engineer")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Python"
    document.save(docx_path)

    parsed = parse_resume(docx_path, "docx", "resume.docx")

    assert parsed.file_type == "docx"
    assert "Alex Engineer" in parsed.normalized_text
    assert "Skill | Python" in parsed.normalized_text


def test_parse_docx_empty_raises(tmp_path: Path) -> None:
    docx_path = tmp_path / "empty.docx"
    document = Document()
    document.save(docx_path)

    with pytest.raises(ResumeParseError, match="no extractable text"):
        parse_resume(docx_path, "docx", "empty.docx")


def test_parse_docx_corrupted_raises(tmp_path: Path) -> None:
    docx_path = tmp_path / "broken.docx"
    docx_path.write_bytes(b"not-a-zip")

    with pytest.raises(ResumeParseError, match="Invalid or corrupted DOCX"):
        parse_resume(docx_path, "docx", "broken.docx")


def test_validation_rejects_unsupported_extension() -> None:
    with pytest.raises(ResumeValidationError, match="Unsupported file type"):
        validate_resume_upload("resume.txt", "text/plain", b"hello", 1024)


def test_validation_rejects_oversized_file() -> None:
    payload = b"a" * 9
    with pytest.raises(ResumeValidationError, match="maximum allowed size"):
        validate_resume_upload("resume.pdf", "application/pdf", payload, 4)


def test_validation_rejects_empty_file() -> None:
    with pytest.raises(ResumeValidationError, match="empty"):
        validate_resume_upload("resume.pdf", "application/pdf", b"", 1024)


def test_validation_rejects_unsafe_filename() -> None:
    with pytest.raises(ResumeValidationError, match="Unsafe filename"):
        validate_resume_upload("../../resume.pdf", "application/pdf", b"%PDF-1.4", 1024)


def test_validation_rejects_invalid_signature() -> None:
    with pytest.raises(ResumeValidationError, match="Invalid PDF file signature"):
        validate_resume_upload("resume.pdf", "application/pdf", b"not-a-pdf", 1024)


def test_normalization_normalizes_line_endings_and_whitespace() -> None:
    raw = "Line 1\r\n\r\n\r\nLine 2    with   spaces\rLine 3\t\t"
    normalized = normalize_resume_text(raw)

    assert normalized == "Line 1\n\nLine 2 with spaces\nLine 3"
