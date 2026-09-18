from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ResumeParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedResume:
    filename: str
    file_type: str
    raw_text: str
    normalized_text: str
    page_count: int


def normalize_resume_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = "\n".join(line.rstrip() for line in normalized.split("\n"))
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    normalized = re.sub(r"[ \t]{2,}", " ", normalized)
    return normalized.strip()


def parse_resume(file_path: Path, file_type: str, filename: str) -> ParsedResume:
    suffix = file_type.lower().lstrip(".")
    if suffix == "pdf":
        raw_text, page_count = _parse_pdf(file_path)
    elif suffix == "docx":
        raw_text, page_count = _parse_docx(file_path)
    else:
        raise ResumeParseError("Unsupported resume file type")

    normalized_text = normalize_resume_text(raw_text)
    if not normalized_text:
        raise ResumeParseError("Uploaded document has no extractable text")

    return ParsedResume(
        filename=filename,
        file_type=suffix,
        raw_text=raw_text,
        normalized_text=normalized_text,
        page_count=page_count,
    )


def _parse_pdf(file_path: Path) -> tuple[str, int]:
    try:
        reader = PdfReader(str(file_path))
    except Exception as exc:  # pragma: no cover - parser exception types vary
        raise ResumeParseError("Invalid or corrupted PDF file") from exc

    sections: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(f"[Page {index}]\n{text.strip()}")

    raw_text = "\n\n".join(sections)
    return raw_text, len(reader.pages)


def _parse_docx(file_path: Path) -> tuple[str, int]:
    try:
        document = Document(str(file_path))
    except Exception as exc:  # pragma: no cover - parser exception types vary
        raise ResumeParseError("Invalid or corrupted DOCX file") from exc

    sections: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(text)

    for table in document.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                sections.append(" | ".join(row_cells))

    raw_text = "\n\n".join(sections)
    return raw_text, max(1, len(document.sections))
